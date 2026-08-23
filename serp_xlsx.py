#!/usr/bin/env python3
"""Сравнение двух срезов Google News и выгрузка в XLSX."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import google_news_parser as gnp

HEADER_FILL = PatternFill("solid", fgColor="1C1814")
HEADER_FONT = Font(color="F6F1E8", bold=True, name="Calibri", size=11)
NEW_FILL = PatternFill("solid", fgColor="D9EAD3")
GONE_FILL = PatternFill("solid", fgColor="F4CCCC")
UP_FILL = PatternFill("solid", fgColor="C9DAF8")
DOWN_FILL = PatternFill("solid", fgColor="FFF2CC")
SAME_FILL = PatternFill("solid", fgColor="F3F3F3")
THIN = Border(
    left=Side(style="thin", color="D7CBB8"),
    right=Side(style="thin", color="D7CBB8"),
    top=Side(style="thin", color="D7CBB8"),
    bottom=Side(style="thin", color="D7CBB8"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_rows(path: Path) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError(f"Ожидался список карточек в {path}")
    return [dict(item) for item in data]


def position_number(value: Any) -> int | None:
    try:
        return int(str(value).split()[0])
    except (TypeError, ValueError):
        return None


def row_key(item: dict[str, Any]) -> tuple[str, str]:
    return (str(item.get("query") or ""), gnp.canonical_url(str(item.get("url") or "")))


def fetched_label(rows: list[dict[str, Any]], fallback: str = "") -> str:
    for item in rows:
        if item.get("fetched_at"):
            return str(item["fetched_at"])
    return fallback


def compare_rows(previous: list[dict[str, Any]], current: list[dict[str, Any]]) -> dict[str, Any]:
    prev_map = {row_key(item): item for item in previous if row_key(item)[1]}
    curr_map = {row_key(item): item for item in current if row_key(item)[1]}
    shared_keys = sorted(set(prev_map) & set(curr_map))
    new_keys = sorted(set(curr_map) - set(prev_map))
    gone_keys = sorted(set(prev_map) - set(curr_map))

    stayed: list[dict[str, Any]] = []
    rose = fell = unchanged = 0
    for key in shared_keys:
        prev_item = prev_map[key]
        curr_item = curr_map[key]
        prev_pos = position_number(prev_item.get("position"))
        curr_pos = position_number(curr_item.get("position"))
        delta = None
        if prev_pos is not None and curr_pos is not None:
            delta = curr_pos - prev_pos
            if delta < 0:
                rose += 1
                change = "поднялась"
            elif delta > 0:
                fell += 1
                change = "опустилась"
            else:
                unchanged += 1
                change = "без изменений"
        else:
            change = "без изменений"
            unchanged += 1
        stayed.append(
            {
                "query": key[0],
                "url": curr_item.get("url") or prev_item.get("url"),
                "canonical_url": key[1],
                "title": curr_item.get("title"),
                "source": curr_item.get("source"),
                "domain": curr_item.get("domain"),
                "position_prev": prev_item.get("position"),
                "position_curr": curr_item.get("position"),
                "delta": delta,
                "change": change,
                "published_at": curr_item.get("published_at"),
                "title_changed": (prev_item.get("title") or "") != (curr_item.get("title") or ""),
                "title_prev": prev_item.get("title"),
            }
        )

    by_query: list[dict[str, Any]] = []
    queries = sorted({item.get("query") for item in previous + current if item.get("query")})
    for query in queries:
        prev_q = [item for item in previous if item.get("query") == query]
        curr_q = [item for item in current if item.get("query") == query]
        prev_urls = {gnp.canonical_url(str(item.get("url") or "")) for item in prev_q}
        curr_urls = {gnp.canonical_url(str(item.get("url") or "")) for item in curr_q}
        prev_urls.discard("")
        curr_urls.discard("")
        top_prev = next((item.get("title") for item in prev_q if str(item.get("position")) == "1"), "")
        top_curr = next((item.get("title") for item in curr_q if str(item.get("position")) == "1"), "")
        by_query.append(
            {
                "query": query,
                "prev_count": len(prev_q),
                "curr_count": len(curr_q),
                "overlap": len(prev_urls & curr_urls),
                "new": len(curr_urls - prev_urls),
                "gone": len(prev_urls - curr_urls),
                "top1_same": bool(top_prev and top_prev == top_curr),
                "top1_prev": top_prev,
                "top1_curr": top_curr,
            }
        )

    prev_domains = Counter(item.get("domain") or "" for item in previous)
    curr_domains = Counter(item.get("domain") or "" for item in current)
    domains = sorted(set(prev_domains) | set(curr_domains))
    domain_rows = [
        {
            "domain": domain,
            "prev": prev_domains.get(domain, 0),
            "curr": curr_domains.get(domain, 0),
            "delta": curr_domains.get(domain, 0) - prev_domains.get(domain, 0),
        }
        for domain in domains
        if domain
    ]
    domain_rows.sort(key=lambda item: (-abs(item["delta"]), -item["curr"], item["domain"]))

    return {
        "previous_at": fetched_label(previous),
        "current_at": fetched_label(current),
        "prev_count": len(previous),
        "curr_count": len(current),
        "new_count": len(new_keys),
        "gone_count": len(gone_keys),
        "stable_count": len(shared_keys),
        "rose": rose,
        "fell": fell,
        "unchanged": unchanged,
        "stayed": stayed,
        "new_rows": [curr_map[key] for key in new_keys],
        "gone_rows": [prev_map[key] for key in gone_keys],
        "by_query": by_query,
        "domains": domain_rows,
    }


def _style_header(sheet: Worksheet, width: int) -> None:
    for col in range(1, width + 1):
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24


def _write_rows(sheet: Worksheet, headers: list[str], rows: list[list[Any]], fills: list[PatternFill | None] | None = None) -> None:
    sheet.append(headers)
    for index, row in enumerate(rows):
        sheet.append(row)
        fill = fills[index] if fills else None
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(index + 2, col)
            cell.alignment = WRAP
            cell.border = THIN
            if fill is not None:
                cell.fill = fill
    _style_header(sheet, len(headers))
    for col, header in enumerate(headers, start=1):
        letter = get_column_letter(col)
        sample = max([len(str(header))] + [len(str(row[col - 1])) for row in rows[:40]], default=10)
        sheet.column_dimensions[letter].width = min(max(sample + 2, 12), 60)


def _card_row(item: dict[str, Any]) -> list[Any]:
    return [
        item.get("query"),
        item.get("position"),
        item.get("title"),
        item.get("source"),
        item.get("published_at"),
        item.get("snippet"),
        item.get("url"),
        item.get("domain"),
        item.get("result_type"),
        item.get("relevance_comment"),
    ]


CARD_HEADERS = ["Запрос", "Позиция", "Заголовок", "СМИ", "Дата", "Сниппет", "URL", "Домен", "Тип", "Релевантность"]


def write_comparison_xlsx(
    previous: list[dict[str, Any]],
    current: list[dict[str, Any]],
    path: Path,
    previous_label: str = "предыдущий скрининг",
    current_label: str = "актуальный срез",
) -> Path:
    cmp = compare_rows(previous, current)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()

    summary = book.active
    summary.title = "Сводка"
    summary_rows = [
        ["Показатель", "Значение"],
        ["Предыдущий скрининг", cmp["previous_at"] or previous_label],
        ["Актуальный срез", cmp["current_at"] or current_label],
        ["Карточек было", cmp["prev_count"]],
        ["Карточек стало", cmp["curr_count"]],
        ["Сохранились (тот же запрос + URL)", cmp["stable_count"]],
        ["Новые публикации", cmp["new_count"]],
        ["Выпали из выдачи", cmp["gone_count"]],
        ["Поднялись в выдаче", cmp["rose"]],
        ["Опустились в выдаче", cmp["fell"]],
        ["Позиция не изменилась", cmp["unchanged"]],
    ]
    for row in summary_rows:
        summary.append(row)
    _style_header(summary, 2)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 56
    summary["A13"] = "Как читать"
    summary["A14"] = (
        "Сравнение идёт по паре «запрос + канонический URL». "
        "Отрицательная дельта позиции значит, что карточка поднялась (например, с 5 на 2 = −3). "
        "Разные публикации одного СМИ не сливаются."
    )
    summary.merge_cells("A14:B16")
    summary["A14"].alignment = WRAP

    stayed_rows = [
        [
            item["query"],
            item["position_prev"],
            item["position_curr"],
            item["delta"],
            item["change"],
            item["title"],
            item["source"],
            item["domain"],
            item["url"],
            "да" if item["title_changed"] else "нет",
            item["published_at"],
        ]
        for item in cmp["stayed"]
    ]
    stayed_fills = []
    for item in cmp["stayed"]:
        if item["change"] == "поднялась":
            stayed_fills.append(UP_FILL)
        elif item["change"] == "опустилась":
            stayed_fills.append(DOWN_FILL)
        else:
            stayed_fills.append(SAME_FILL)
    stayed_sheet = book.create_sheet("Сравнение позиций")
    _write_rows(
        stayed_sheet,
        [
            "Запрос",
            "Позиция было",
            "Позиция стало",
            "Дельта",
            "Изменение",
            "Заголовок",
            "СМИ",
            "Домен",
            "URL",
            "Заголовок изменился",
            "Дата публикации",
        ],
        stayed_rows,
        stayed_fills,
    )

    new_sheet = book.create_sheet("Новые")
    _write_rows(new_sheet, CARD_HEADERS, [_card_row(item) for item in cmp["new_rows"]], [NEW_FILL] * len(cmp["new_rows"]))

    gone_sheet = book.create_sheet("Выпавшие")
    _write_rows(gone_sheet, CARD_HEADERS, [_card_row(item) for item in cmp["gone_rows"]], [GONE_FILL] * len(cmp["gone_rows"]))

    current_sheet = book.create_sheet("Актуальный срез")
    _write_rows(current_sheet, CARD_HEADERS, [_card_row(item) for item in current])

    previous_sheet = book.create_sheet("Предыдущий скрининг")
    _write_rows(previous_sheet, CARD_HEADERS, [_card_row(item) for item in previous])

    query_sheet = book.create_sheet("По запросам")
    _write_rows(
        query_sheet,
        [
            "Запрос",
            "Было",
            "Стало",
            "Пересечение URL",
            "Новых",
            "Выпало",
            "Топ-1 тот же",
            "Топ-1 было",
            "Топ-1 стало",
        ],
        [
            [
                item["query"],
                item["prev_count"],
                item["curr_count"],
                item["overlap"],
                item["new"],
                item["gone"],
                "да" if item["top1_same"] else "нет",
                item["top1_prev"],
                item["top1_curr"],
            ]
            for item in cmp["by_query"]
        ],
    )

    domain_sheet = book.create_sheet("СМИ и домены")
    _write_rows(
        domain_sheet,
        ["Домен", "Было карточек", "Стало карточек", "Дельта"],
        [[item["domain"], item["prev"], item["curr"], item["delta"]] for item in cmp["domains"]],
    )

    book.save(path)
    return path


def write_current_xlsx(rows: list[dict[str, Any]], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Выдача"
    _write_rows(sheet, CARD_HEADERS, [_card_row(item) for item in rows])
    book.save(path)
    return path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Сравнение срезов Google News в XLSX")
    parser.add_argument("--previous", required=True, help="JSON предыдущего скрининга")
    parser.add_argument("--current", required=True, help="JSON актуального среза")
    parser.add_argument("--output", required=True, help="Путь к XLSX")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    previous = load_rows(Path(args.previous))
    current = load_rows(Path(args.current))
    path = write_comparison_xlsx(previous, current, Path(args.output))
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
