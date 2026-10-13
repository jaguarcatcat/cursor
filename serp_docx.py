#!/usr/bin/env python3
"""Выгрузка среза и сравнения Google News в DOCX."""

from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import google_news_parser as gnp
import serp_xlsx as sx


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell(cell, text: Any, *, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("" if text is None else str(text))
    run.font.size = Pt(9)
    run.font.name = "Calibri"
    run.bold = bold
    if fill:
        _shade_cell(cell, fill)


def _add_table(document: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        _set_cell(table.rows[0].cells[index], header, bold=True, fill="1C1814")
        for paragraph in table.rows[0].cells[index].paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0xF6, 0xF1, 0xE8)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            _set_cell(table.rows[row_index].cells[col_index], value)
    document.add_paragraph()


def _heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT


def write_current_docx(
    rows: list[dict[str, Any]],
    path: Path,
    checked_at: str = "",
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    _heading(document, "Актуальный срез Google Новости")
    meta = document.add_paragraph()
    meta.add_run(checked_at or sx.fetched_label(rows) or "").italic = True
    document.add_paragraph(
        "Источник: Google News / Google Новости, регион RU, язык ru. "
        "Органический поиск не использовался. Сниппеты не выдумывались."
    )
    by_query: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in rows:
        by_query[str(item.get("query") or "")].append(item)
    for query, items in by_query.items():
        items = sorted(items, key=lambda row: sx.position_number(row.get("position")) or 999)
        _heading(document, f"Запрос «{query}»", 2)
        _add_table(
            document,
            ["#", "СМИ", "Заголовок", "Дата", "Домен", "URL"],
            [
                [
                    item.get("position"),
                    item.get("source"),
                    item.get("title"),
                    item.get("published_at"),
                    item.get("domain"),
                    item.get("url"),
                ]
                for item in items
            ],
        )
    _heading(document, "Аналитика", 2)
    tone = Counter(
        gnp.classify_tone(str(item.get("title") or ""), str(item.get("url") or ""), str(item.get("snippet") or ""))
        for item in rows
    )
    topic = Counter(gnp.classify_topic(str(item.get("title") or ""), str(item.get("url") or "")) for item in rows)
    source = Counter(str(item.get("source") or "") for item in rows)
    document.add_paragraph("Тональность: " + "; ".join(f"{name} ({count})" for name, count in tone.most_common()))
    document.add_paragraph("Темы: " + "; ".join(f"{name} ({count})" for name, count in topic.most_common(8)))
    document.add_paragraph("СМИ: " + "; ".join(f"{name} ({count})" for name, count in source.most_common(8)))
    document.save(path)
    return path


def write_comparison_docx(
    previous: list[dict[str, Any]],
    current: list[dict[str, Any]],
    path: Path,
    previous_label: str = "предыдущий скрининг",
    current_label: str = "актуальный срез",
    baseline: list[dict[str, Any]] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    current = sx.enrich_publisher_urls(current, previous, *([baseline] if baseline else []))
    previous = sx.enrich_publisher_urls(previous, *([baseline] if baseline else []))
    cmp = sx.compare_rows(previous, current)
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    _heading(document, "Сравнение выдачи Google Новости")
    document.add_paragraph(f"{previous_label} → {current_label}")
    _heading(document, "Сводка", 2)
    _add_table(
        document,
        ["Показатель", "Значение"],
        [
            ["Предыдущий скрининг", cmp["previous_at"] or previous_label],
            ["Актуальный срез", cmp["current_at"] or current_label],
            ["Карточек было", cmp["prev_count"]],
            ["Карточек стало", cmp["curr_count"]],
            ["Сохранились", cmp["stable_count"]],
            ["Новые", cmp["new_count"]],
            ["Выпали", cmp["gone_count"]],
            ["Поднялись", cmp["rose"]],
            ["Опустились", cmp["fell"]],
            ["Без сдвига", cmp["unchanged"]],
        ],
    )
    _heading(document, "По запросам", 2)
    _add_table(
        document,
        ["Запрос", "Было", "Стало", "Пересечение", "Новых", "Выпало", "Топ-1 тот же", "Топ-1 было", "Топ-1 стало"],
        [
            [
                item["query"],
                item["prev_count"],
                item["curr_count"],
                item["overlap"],
                item["new"],
                item["gone"],
                "да" if item["top1_same"] else "нет",
                item["top1_prev"],
                item["top1_curr"],
            ]
            for item in cmp["by_query"]
        ],
    )
    _heading(document, "Сводная таблица сравнения", 2)
    table_rows, _fills = sx.build_summary_table(cmp)
    _add_table(
        document,
        [
            "Запрос",
            "Заголовок",
            "СМИ",
            "Домен",
            f"Позиция: {previous_label}",
            f"Позиция: {current_label}",
            "Дельта",
            "Статус",
            "URL",
        ],
        table_rows,
    )
    if cmp["new_rows"]:
        _heading(document, "Новые карточки", 2)
        _add_table(
            document,
            ["Запрос", "#", "СМИ", "Заголовок", "URL"],
            [
                [item.get("query"), item.get("position"), item.get("source"), item.get("title"), item.get("url")]
                for item in cmp["new_rows"]
            ],
        )
    if cmp["gone_rows"]:
        _heading(document, "Выпавшие карточки", 2)
        _add_table(
            document,
            ["Запрос", "#", "СМИ", "Заголовок", "URL"],
            [
                [item.get("query"), item.get("position"), item.get("source"), item.get("title"), item.get("url")]
                for item in cmp["gone_rows"]
            ],
        )
    _heading(document, "Аналитика", 2)
    for block, conclusion in sx.build_analytics_lines(previous, current, cmp)[1:]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{block}: ")
        run.bold = True
        paragraph.add_run(str(conclusion))
    document.save(path)
    return path
