#!/usr/bin/env python3
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import serp_xlsx


PREV = [
    {
        "query": "лобов вадим",
        "position": "1",
        "title": "Старый топ",
        "source": "Sostav.ru",
        "url": "https://www.sostav.ru/blogs/1",
        "published_at": "2026-07-13 00:04 MSK",
        "snippet": "сниппет Google недоступен",
        "domain": "sostav.ru",
        "result_type": "новость",
        "relevance_comment": "высокая",
        "fetched_at": "2026-08-17T02:09:26+03:00",
    },
    {
        "query": "лобов вадим",
        "position": "2",
        "title": "Выпадет",
        "source": "РЕПОСТ",
        "url": "https://repost.news/old",
        "published_at": "2024-07-12 15:04 MSK",
        "snippet": "сниппет Google недоступен",
        "domain": "repost.news",
        "result_type": "новость",
        "relevance_comment": "высокая",
        "fetched_at": "2026-08-17T02:09:26+03:00",
    },
]

CURR = [
    {
        "query": "лобов вадим",
        "position": "1",
        "title": "Новый топ",
        "source": "Forbes.ru",
        "url": "https://www.forbes.ru/new",
        "published_at": "2026-08-20 11:00 MSK",
        "snippet": "сниппет Google недоступен",
        "domain": "forbes.ru",
        "result_type": "новость",
        "relevance_comment": "высокая",
        "fetched_at": "2026-08-23T19:00:00+03:00",
    },
    {
        "query": "лобов вадим",
        "position": "3",
        "title": "Старый топ",
        "source": "Sostav.ru",
        "url": "http://sostav.ru/blogs/1?utm_source=x",
        "published_at": "2026-07-13 00:04 MSK",
        "snippet": "сниппет Google недоступен",
        "domain": "sostav.ru",
        "result_type": "новость",
        "relevance_comment": "высокая",
        "fetched_at": "2026-08-23T19:00:00+03:00",
    },
]


class CompareTests(unittest.TestCase):
    def test_compare_detects_new_gone_and_drop(self) -> None:
        cmp = serp_xlsx.compare_rows(PREV, CURR)
        self.assertEqual(cmp["new_count"], 1)
        self.assertEqual(cmp["gone_count"], 1)
        self.assertEqual(cmp["stable_count"], 1)
        self.assertEqual(cmp["stayed"][0]["delta"], 2)
        self.assertEqual(cmp["stayed"][0]["change"], "опустилась")
        self.assertEqual(cmp["new_rows"][0]["domain"], "forbes.ru")
        self.assertEqual(cmp["gone_rows"][0]["domain"], "repost.news")

    def test_xlsx_has_analysis_sheets(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cmp.xlsx"
            serp_xlsx.write_comparison_xlsx(PREV, CURR, path)
            book = load_workbook(path)
            self.assertIn("Сводка", book.sheetnames)
            self.assertIn("Сравнение позиций", book.sheetnames)
            self.assertIn("Новые", book.sheetnames)
            self.assertIn("Выпавшие", book.sheetnames)
            self.assertIn("Актуальный срез", book.sheetnames)
            self.assertIn("Предыдущий скрининг", book.sheetnames)
            self.assertEqual(book["Новые"].max_row, 2)
            self.assertEqual(book["Выпавшие"].max_row, 2)
            self.assertEqual(book["Актуальный срез"].max_row, 3)


if __name__ == "__main__":
    unittest.main()
